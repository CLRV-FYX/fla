# -*- coding: utf-8 -*-
"""
FLA v1.13 - OMML 公式预处理
背景: LibreOffice 无法导入 PPTX 中的 OMML 公式(PowerPoint 公式编辑器),
     转换时会被静默丢弃 → PDF 预览/放映背景层全部缺公式。
方案: 转换前把含公式的文本框渲染为高清 PNG(matplotlib mathtext),
     原位替换并保留 shape_id → 动画时序(spTgt spid)依然匹配。
公式图片天然可参与放映引擎的入场动画(fade/fly/zoom/wipe)。
失败策略: 单个形状替换失败 → 保留原状(pptx_anim 的 OMML 守卫会让该页回退静态);
         整体失败 → 返回 None, 转换流程照旧(不比现状更差)。
"""
import io
import re
from pathlib import Path

from lxml import etree

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
A14 = "http://schemas.microsoft.com/office/drawing/2010/main"

EMU_PER_IN = 914400
DPI = 300

_fonts_ready = False


def _register_fonts():
    """把系统/用户目录的中文字体注册进 matplotlib(其自带缓存不含后装字体)."""
    global _fonts_ready
    if _fonts_ready:
        return
    _fonts_ready = True
    try:
        import matplotlib
        matplotlib.use("Agg")
        from matplotlib import font_manager
        import glob
        import os
        dirs = ["/usr/share/fonts", os.path.expanduser("~/.fonts"),
                "/var/www/fonts", "/usr/local/share/fonts"]
        pats = set()
        for d in dirs:
            if d and os.path.isdir(d):
                for root, _, files in os.walk(d):
                    for f in files:
                        if f.lower().endswith((".otf", ".ttf")):
                            pats.add(os.path.join(root, f))
        for p in pats:
            try:
                font_manager.fontManager.addfont(p)
            except Exception:
                pass
    except Exception:
        pass


def _preferred_family():
    """按优先级挑选已注册的中文字体族."""
    try:
        from matplotlib import font_manager
        names = {f.name for f in font_manager.fontManager.ttflist}
        for cand in ["Noto Sans CJK SC", "Noto Sans CJK JP", "LXGW WenKai",
                     "WenQuanYi Zen Hei", "Source Han Sans SC", "AR PL UMing CN"]:
            if cand in names:
                return cand
    except Exception:
        pass
    return "DejaVu Sans"


# ---------------- OMML → LaTeX ----------------

_CHAR_MAP = {
    "−": "-", "×": r"\times ", "÷": r"\div ", "±": r"\pm ", "⋅": r"\cdot ",
    "°": r"^{\circ}", "′": "'", "″": "''",
    "…": r"\ldots ", "⋯": r"\cdots ", "∞": r"\infty ",
    "→": r"\to ", "←": r"\leftarrow ", "⇒": r"\Rightarrow ", "⇔": r"\Leftrightarrow ",
    "≤": r"\leq ", "≥": r"\geq ", "≠": r"\neq ", "≈": r"\approx ",
    "≅": r"\cong ", "∼": r"\sim ", "∽": r"\sim ", "∝": r"\propto ",
    "∠": r"\angle ", "⊥": r"\perp ", "∥": r"\parallel ", "△": r"\triangle ",
    "∈": r"\in ", "∉": r"\notin ", "⊆": r"\subseteq ", "⊂": r"\subset ",
    "⊇": r"\supseteq ", "⊃": r"\supset ", "∪": r"\cup ", "∩": r"\cap ",
    "∅": r"\emptyset ", "⊙": r"\odot ", "⊕": r"\oplus ", "⊗": r"\otimes ",
    "∑": r"\sum ", "∫": r"\int ", "∏": r"\prod ",
    "²": "^{2}", "³": "^{3}", "ⁿ": "^{n}", "½": r"\frac{1}{2}",
}
_FUNC_MAP = {"lim", "sin", "cos", "tan", "cot", "sec", "csc", "log", "ln",
             "exp", "max", "min", "det", "gcd", "sup", "inf", "arg", "dim"}
_LATEX_ESC = {"{": r"\{", "}": r"\}", "_": r"\_", "&": r"\&", "$": r"\$",
              "#": r"\#", "%": r"\%", "~": r"\sim ", "\\": r"\backslash "}


def _mt_text(el):
    """m:t 文本 → latex 片段."""
    t = el.text or ""
    out = []
    for ch in t:
        if ch in _CHAR_MAP:
            out.append(_CHAR_MAP[ch])
        elif ch in _LATEX_ESC:
            out.append(_LATEX_ESC[ch])
        elif ch == "^":
            out.append(r"\hat{\;}")
        else:
            out.append(ch)
    return "".join(out)


def _kids(el, tag):
    return [c for c in el if c.tag == "{%s}%s" % (M, tag)]


def _first(el, tag):
    k = _kids(el, tag)
    return k[0] if k else None


def omml_to_latex(el):
    """递归转换 OMML 元素为 matplotlib mathtext 可渲染的 latex."""
    tag = el.tag.split("}")[1] if "}" in el.tag else el.tag
    if tag in ("oMath", "oMathPara", "oMathParaPr"):
        return "".join(omml_to_latex(c) for c in el if c.tag.split("}")[1] not in ("oMathParaPr",))
    if tag == "r":
        return "".join(_mt_text(t) for t in el.iter("{%s}t" % M))
    if tag == "f":
        num = omml_to_latex(_first(el, "num")) if _first(el, "num") is not None else ""
        den = omml_to_latex(_first(el, "den")) if _first(el, "den") is not None else ""
        return r"\frac{%s}{%s}" % (num, den)
    if tag == "rad":
        pr = _first(el, "radPr")
        hide = pr is not None and _first(pr, "degHide") is not None
        deg = _first(el, "deg")
        e = _first(el, "e")
        body = omml_to_latex(e) if e is not None else ""
        d = ""
        if deg is not None and not hide:
            d = omml_to_latex(deg)
        return (r"\sqrt[%s]{%s}" % (d, body)) if d else (r"\sqrt{%s}" % body)
    if tag in ("sSup", "sSub", "sSubSup"):
        e = omml_to_latex(_first(el, "e")) if _first(el, "e") is not None else ""
        sup = _first(el, "sup")
        sub = _first(el, "sub")
        out = "{" + e + "}"
        if sub is not None:
            out += "_{" + omml_to_latex(sub) + "}"
        if sup is not None:
            out += "^{" + omml_to_latex(sup) + "}"
        return out
    if tag == "sPre":
        e = omml_to_latex(_first(el, "e")) if _first(el, "e") is not None else ""
        sup, sub = _first(el, "sup"), _first(el, "sub")
        out = "{}"
        if sub is not None:
            out += "_{" + omml_to_latex(sub) + "}"
        if sup is not None:
            out += "^{" + omml_to_latex(sup) + "}"
        return out + "{" + e + "}"
    if tag == "nary":
        pr = _first(el, "naryPr")
        chr_ = "∫"
        if pr is not None:
            c = _first(pr, "chr")
            if c is not None and c.get("{%s}val" % M):
                chr_ = c.get("{%s}val" % M)
        op = {"∑": r"\sum ", "∏": r"\prod ", "∫": r"\int ", "∬": r"\iint ",
              "∮": r"\oint ", "∐": r"\coprod "}.get(chr_, chr_ + " ")
        sub, sup, e = _first(el, "sub"), _first(el, "sup"), _first(el, "e")
        out = op
        if sub is not None:
            out += "_{" + omml_to_latex(sub) + "}"
        if sup is not None:
            out += "^{" + omml_to_latex(sup) + "}"
        if e is not None:
            out += " " + omml_to_latex(e)
        return out
    if tag == "d":
        pr = _first(el, "dPr")
        beg, end = "(", ")"
        if pr is not None:
            b, e2 = _first(pr, "begChr"), _first(pr, "endChr")
            if b is not None:
                beg = b.get("{%s}val" % M) or ""
            if e2 is not None:
                end = e2.get("{%s}val" % M) or ""
        bm = {"(": r"\left(", "[": r"\left[", "{": r"\left\{", "|": r"\left|",
              "": r"\left.", "⟨": r"\langle "}.get(beg, beg)
        em = {")": r"\right)", "]": r"\right]", "}": r"\right\}", "|": r"\right|",
              "": r"\right.", "⟩": r"\rangle "}.get(end, end)
        inner = "".join(omml_to_latex(x) for x in _kids(el, "e"))
        return bm + inner + em
    if tag == "func":
        nm = _first(el, "fName")
        name = omml_to_latex(nm).strip() if nm is not None else ""
        if name in _FUNC_MAP:
            name = "\\" + name + " "
        e = _first(el, "e")
        return name + (omml_to_latex(e) if e is not None else "")
    if tag == "limLow":
        e, lim = _first(el, "e"), _first(el, "lim")
        base = omml_to_latex(e) if e is not None else ""
        l = omml_to_latex(lim) if lim is not None else ""
        if base.strip() in ("lim", "max", "min"):
            return "\\" + base.strip() + "_{" + l + "} "
        return "{" + base + "}_{" + l + "}"
    if tag == "limUpp":
        e, lim = _first(el, "e"), _first(el, "lim")
        base = omml_to_latex(e) if e is not None else ""
        l = omml_to_latex(lim) if lim is not None else ""
        return "{" + base + "}^{" + l + "}"
    if tag == "bar":
        e = _first(el, "e")
        return r"\bar{%s}" % (omml_to_latex(e) if e is not None else "")
    if tag == "acc":
        pr = _first(el, "accPr")
        chr_ = "̂"
        if pr is not None and _first(pr, "chr") is not None:
            chr_ = _first(pr, "chr").get("{%s}val" % M) or "̂"
        cmd = {"̂": "hat", "̃": "tilde", "⃗": "vec", "̄": "bar", "˙": "dot",
               "̈": "ddot", "́": "acute", "̌": "check", "̆": "breve"}.get(chr_, "hat")
        e = _first(el, "e")
        return r"\%s{%s}" % (cmd, omml_to_latex(e) if e is not None else "")
    if tag == "m":  # 矩阵: mathtext 无矩阵环境, 退化为同行排列
        rows = []
        for mr in _kids(el, "mr"):
            rows.append(r" \; ".join(omml_to_latex(x) for x in _kids(mr, "e")))
        return r" \;\; ".join(rows)
    if tag == "eqArr":
        return r" \\\\ ".join(omml_to_latex(x) for x in _kids(el, "e"))
    if tag in ("box", "borderBox", "groupChr", "sSubSupPr", "ctrlPr", "fPr",
               "radPr", "naryPr", "dPr", "funcPr", "limLowPr", "accPr", "barPr", "rPr"):
        inner = "".join(omml_to_latex(c) for c in el if c.tag.split("}")[1] in
                        ("e", "num", "den", "sub", "sup", "deg", "fName", "lim", "r", "oMath"))
        return inner
    if tag in ("num", "den", "e", "sub", "sup", "deg", "lim", "fName", "mr", "oMathParaPr"):
        return "".join(omml_to_latex(c) for c in el)
    # 未知元素: 拼接所有 m:t 文本兜底
    return "".join(_mt_text(t) for t in el.iter("{%s}t" % M))


# ---------------- 渲染 ----------------

def _escape_text(t):
    return t.replace("$", r"\$")


def _render_lines(lines):
    """lines: [(plain_or_latex_text, fontsize_pt, bold, align)] → (png_bytes, emu_w, emu_h)"""
    import matplotlib.pyplot as plt
    _register_fonts()
    fam = _preferred_family()
    plt.rcParams["font.family"] = [fam, "DejaVu Sans"]
    plt.rcParams["mathtext.fontset"] = "dejavusans"
    fig = plt.figure(figsize=(12, 6))
    y = 0.98
    for text, fs, bold, _al in lines:
        fig.text(0.01, y, text, fontsize=fs, va="top", ha="left",
                 fontweight="bold" if bold else "normal")
        y -= min(0.28, (fs * 1.5) / 300)
    buf = io.BytesIO()
    fig.savefig(buf, dpi=DPI, bbox_inches="tight", pad_inches=0.03, transparent=True)
    plt.close(fig)
    png = buf.getvalue()
    from PIL import Image
    im = Image.open(io.BytesIO(png))
    emu_w = int(im.width / DPI * EMU_PER_IN)
    emu_h = int(im.height / DPI * EMU_PER_IN)
    return png, emu_w, emu_h


# ---------------- 形状检测与替换 ----------------

def shape_has_math(sp_el):
    for el in sp_el.iter():
        t = el.tag
        if t == "{%s}oMath" % M or t == "{%s}oMathPara" % M:
            return True
        if t == "{%s}m" % A14:
            return True
    return False


def _para_lines(p_el):
    """a:p → [(display_text, fontsize, bold, align)] 或 None(无公式时返回文本行)."""
    has_math = any(c.tag == "{%s}m" % A14 for c in p_el)
    parts = []
    max_sz = 0.0
    bold = False
    for c in p_el:
        if c.tag == "{%s}r" % A:
            t = "".join(x.text or "" for x in c.iter("{%s}t" % A))
            if not t:
                continue
            rPr = c.find("{%s}rPr" % A)
            sz = 18.0
            if rPr is not None and rPr.get("sz"):
                try:
                    sz = int(rPr.get("sz")) / 100.0
                except ValueError:
                    pass
                if rPr.find("{%s}b" % A) is not None:
                    bold = True
            max_sz = max(max_sz, sz)
            parts.append(("t", _escape_text(t), sz))
        elif c.tag == "{%s}m" % A14:
            latex = ""
            for om in c.iter("{%s}oMath" % M):
                latex += omml_to_latex(om)
            parts.append(("m", latex, 0))
            if max_sz <= 0:
                max_sz = 20.0
        elif c.tag == "{%s}AlternateContent" % MC:
            for om in c.iter("{%s}oMath" % M):
                parts.append(("m", omml_to_latex(om), 0))
                if max_sz <= 0:
                    max_sz = 20.0
    if not parts:
        return None, False
    display = "".join(p[1] if p[0] == "t" else "$" + p[1] + "$" for p in parts)
    pPr = p_el.find("{%s}pPr" % A)
    algn = pPr.get("algn") if pPr is not None else None
    return (display, max_sz or 18.0, bold, algn), has_math


def render_txbody(txBody_el):
    """任意 txBody(形状/表格单元格) → (png, emu_w, emu_h). 无公式返回 None.
    v1.14: 供 pptx_anim 在表格单元格/残留公式形状上直接调用."""
    if txBody_el is None:
        return None
    lines = []
    any_math = False
    for p_el in txBody_el.findall("{%s}p" % A):
        line, hm = _para_lines(p_el)
        if line:
            lines.append(line)
        any_math = any_math or hm
    if not any_math or not lines:
        return None
    return _render_lines(lines)


def preprocess(src, outdir):
    """含 OMML 公式的文本框 → 图片. 返回新 pptx 路径; 无公式/失败返回 None."""
    from pptx import Presentation
    prs = Presentation(str(src))
    replaced = 0
    tmpdir = Path(outdir) / "_omml_png"
    tmpdir.mkdir(parents=True, exist_ok=True)
    def cnvpr(sp):
        for ns in (P, A):
            nv = sp.find(".//{%s}cNvPr" % ns)
            if nv is not None:
                return nv
        return None

    for slide in prs.slides:
        tree = slide.shapes._spTree
        for sp in list(tree):
            nv = cnvpr(sp)
            spid = nv.get("id") if nv is not None else None
            txBody = sp.find("{%s}txBody" % P)
            if txBody is None:
                txBody = sp.find("{%s}txBody" % A)
            if txBody is None or spid is None:
                continue
            if not shape_has_math(sp):
                continue
            # 几何
            xfrm = sp.find(".//{%s}xfrm" % A)
            if xfrm is None:
                continue
            off, ext = xfrm.find("{%s}off" % A), xfrm.find("{%s}ext" % A)
            if off is None or ext is None:
                continue
            try:
                L, T = int(off.get("x")), int(off.get("y"))
                W, H = int(ext.get("cx")), int(ext.get("cy"))
            except (TypeError, ValueError):
                continue
            lines = []
            any_math = False
            for p_el in txBody.findall("{%s}p" % A):
                line, hm = _para_lines(p_el)
                if line:
                    lines.append(line)
                any_math = any_math or hm
            if not any_math or not lines:
                continue
            try:
                png, emu_w, emu_h = _render_lines(lines)
            except Exception:
                continue
            if emu_h > H and H > 0:  # 过高: 压到盒内
                ratio = H / emu_h
                emu_h = int(emu_h * ratio)
                emu_w = int(emu_w * ratio)
            all_center = all((ln[3] == "ctr") for ln in lines)
            left = L + (W - emu_w) // 2 if all_center else L
            top = T + max(0, (H - emu_h) // 2)
            png_path = tmpdir / ("m_%s_%d.png" % (spid, replaced))
            png_path.write_bytes(png)
            idx = list(tree).index(sp)
            old_name = nv.get("name") or ""
            tree.remove(sp)
            pic = slide.shapes.add_picture(str(png_path), left, top,
                                           width=emu_w, height=emu_h)
            pnv = cnvpr(pic._element)
            pnv.set("id", spid)
            pnv.set("name", old_name or ("formula_" + spid))
            pel = pic._element
            tree.remove(pel)
            tree.insert(idx, pel)
            replaced += 1
    if not replaced:
        return None
    out = Path(outdir) / "_omml_pre.pptx"
    prs.save(str(out))
    return out
